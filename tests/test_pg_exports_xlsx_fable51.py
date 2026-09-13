"""The .xlsx export format (migration 035, Stream C, 2026-09-13).

Source-level (no database): the renderer over a synthetic CSV result --
sheet order, frozen header, autofilter, typed cells, the Indian currency
format, the SUM totals over the counted range and nothing else, the
formula-injection guard on every string cell, the Summary / Applied Filters /
Metadata sheets, the filename, and the two refusals (header drift, row-count
drift). The wheels the bundle vendors are named.

Live (`@pytest.mark.pg`, SKIPPED without CAPEX_DB_URL -- a skip is not a
pass): an xlsx job runs through the SAME chunked, digested pipeline as a CSV
job and its result is a workbook whose Data sheet holds exactly the counted
rows; a csv job is unchanged; an unknown format is refused at creation.
"""
from __future__ import annotations

import io
import os
import re
import sys as _sys
import uuid
from datetime import datetime, timezone
from decimal import Decimal
from pathlib import Path as _Path

_TESTS_DIR = _Path(__file__).resolve().parent
if str(_TESTS_DIR) not in _sys.path:
    _sys.path.insert(0, str(_TESTS_DIR))

from conftest_pg import (  # noqa: E402,F401  (re-exported as fixtures)
    pg_admin_connection, pg_connection, pg_database, pg_disposable_db_name,
    pg_scope, pg_template, pg_url,
)
from test_pg_reservations import _seed  # noqa: E402

import pytest  # noqa: E402
from openpyxl import load_workbook  # noqa: E402

from app.backend.pg import exports as export_svc  # noqa: E402
from app.backend.pg import exports_xlsx as xlsx  # noqa: E402
from app.backend.pg import migrate_pg  # noqa: E402
from app.backend.pg import principal_scope  # noqa: E402
from app.backend.pg.engine import Scope  # noqa: E402

ROOT = _Path(__file__).resolve().parents[1]
PG = pytest.mark.skipif(
    not os.environ.get("CAPEX_DB_URL"),
    reason=("PostgreSQL not configured; set CAPEX_DB_URL to run against a live "
            "database. THIS IS A SKIP, NOT A PASS."),
)

COLUMNS = (
    export_svc.Column("label", "x"),
    export_svc.Column("amount", "x", "paise"),
    export_svc.Column("count", "x", "int"),
    export_svc.Column("ratio", "x", "numeric"),
    export_svc.Column("active", "x", "bool"),
    export_svc.Column("on", "x", "date"),
    export_svc.Column("at", "x", "timestamp"),
)
ORDER = [c.name for c in COLUMNS]
CSV = (
    "label,amount,count,ratio,active,on,at\n"
    "plain,1234567.89,3,2.5,true,2026-09-13,2026-09-13T10:00:00+00:00\n"
    "'=cmd|' /c calc'!A0,-0.05,4,0.25,false,2026-01-01,2026-01-01T00:00:00+05:30\n"
    "=HYPERLINK(\"http://x\"),0.00,,,,,\n"
)
WHEN = datetime(2026, 9, 13, 12, 30, tzinfo=timezone.utc)


def _render(**overrides) -> bytes:
    kwargs = dict(dataset_name="synthetic", dataset_title="A synthetic report",
                  columns=COLUMNS, column_order=ORDER, csv_body=CSV,
                  filters={"project_ids": ["P-1", "P-2"], "date_from": "2026-01-01"},
                  requested_by="U-X", export_job_id="EXP-1", rows_total=3,
                  csv_sha256="abc123", chunk_rows=5000, generated_at=WHEN,
                  app_version="test 0.0")
    kwargs.update(overrides)
    return xlsx.render_workbook(**kwargs)


def _book(data: bytes):
    return load_workbook(io.BytesIO(data))


# =========================================================================
# Source-level
# =========================================================================
def test_035_is_discovered_and_widens_the_same_named_check():
    versions = [m.version for m in migrate_pg.discover()]
    assert versions == sorted(versions) and versions.count("035") == 1
    text = (ROOT / "migrations" / "pg" / "035_export_xlsx.sql").read_text(encoding="utf-8")
    assert "ck_export_job_output_format" in text and "'xlsx'" in text
    assert export_svc.OUTPUT_FORMATS == ("csv", "xlsx")
    req = (ROOT / "requirements.txt").read_text(encoding="utf-8")
    assert "openpyxl==3.1.5" in req


def test_the_workbook_has_the_four_sheets_a_frozen_header_and_an_autofilter():
    wb = _book(_render())
    assert wb.sheetnames == ["Summary", "Data", "Applied Filters", "Metadata"]
    data = wb["Data"]
    assert data.freeze_panes == "A2"
    assert data.auto_filter.ref == "A1:G4"
    assert [c.value for c in data[1]] == ORDER
    assert data["A1"].font.bold
    assert data.column_dimensions["A"].width >= 8


def test_every_cell_is_typed_from_its_kind_and_money_carries_the_indian_format():
    data = _book(_render())["Data"]
    amount = data["B2"]
    assert amount.data_type == "n" and Decimal(str(amount.value)) == Decimal("1234567.89")
    assert amount.number_format == xlsx.INDIAN_CURRENCY_FORMAT
    assert Decimal(str(data["B3"].value)) == Decimal("-0.05")
    assert data["C2"].value == 3 and data["C2"].number_format == xlsx.INT_FORMAT
    assert Decimal(str(data["D2"].value)) == Decimal("2.5")
    assert data["E2"].value is True and data["E3"].value is False
    assert data["F2"].value.date() == datetime(2026, 9, 13).date()
    assert data["F2"].number_format == xlsx.DATE_FORMAT
    assert data["G2"].value == datetime(2026, 9, 13, 10, 0)
    assert data["G3"].value == datetime(2025, 12, 31, 18, 30), "IST rendered as UTC"
    assert data["G2"].number_format == xlsx.TIMESTAMP_FORMAT
    assert data["C4"].value in (None, ""), "an empty CSV cell stays empty"


def test_the_formula_injection_guard_holds_on_every_string_cell():
    data = _book(_render())["Data"]
    guarded = data["A3"]
    assert guarded.data_type == "s" and guarded.value.startswith("'=cmd")
    raw = data["A4"]
    assert raw.data_type == "s" and raw.value == '=HYPERLINK("http://x")', (
        "a value beginning with = that reached the CSV unguarded is still a STRING cell")
    # No formula anywhere but the totals row.
    for row in data.iter_rows(min_row=1, max_row=4):
        for cell in row:
            assert cell.data_type != "f", cell.coordinate


def test_the_totals_row_sums_the_counted_range_and_only_the_numeric_kinds():
    data = _book(_render())["Data"]
    assert data["A5"].value == "Total"
    assert data["B5"].value == "=SUM(B2:B4)" and data["B5"].number_format == xlsx.INDIAN_CURRENCY_FORMAT
    assert data["C5"].value == "=SUM(C2:C4)"
    assert data["D5"].value == "=SUM(D2:D4)"
    for col in ("E5", "F5", "G5"):
        assert data[col].value is None, col
    assert data.max_row == 5
    for row in data.iter_rows(min_row=5, max_row=5):
        for cell in row:
            if cell.data_type == "f":
                assert re.fullmatch(r"=SUM\([A-Z]+2:[A-Z]+4\)", cell.value), cell.value


def test_summary_filters_and_metadata_say_what_the_file_is():
    wb = _book(_render())
    summary = {row[0].value: row[1].value for row in wb["Summary"].iter_rows()}
    assert summary["Report"] == "A synthetic report" and summary["Dataset"] == "synthetic"
    assert summary["Rows"] == 3 and summary["Filters applied"] == 2
    assert summary["CSV result SHA-256"] == "abc123"
    assert summary["Export job"] == "EXP-1" and summary["Application"] == "test 0.0"
    assert summary["Generated at (UTC)"] == "2026-09-13T12:30:00+00:00"
    applied = [(r[0].value, r[1].value) for r in wb["Applied Filters"].iter_rows(min_row=2)]
    assert applied == [("date_from", "2026-01-01"), ("project_ids", "P-1, P-2")]
    unfiltered = _book(_render(filters={}))["Applied Filters"]
    assert unfiltered["A2"].value == "none"
    meta = {r[0].value: (r[1].value, r[2].value)
            for r in wb["Metadata"].iter_rows(min_row=2, max_row=8)}
    assert meta["amount"] == ("paise", xlsx.INDIAN_CURRENCY_FORMAT)
    assert meta["label"] == ("text", "text")


def test_the_renderer_refuses_header_drift_and_row_count_drift():
    with pytest.raises(ValueError):
        _render(column_order=list(reversed(ORDER)))
    with pytest.raises(ValueError):
        _render(rows_total=2)
    assert xlsx.xlsx_filename("budget_ledger_cells", WHEN) == "budget_ledger_cells-2026-09-13.xlsx"


def test_the_bundle_wheels_carry_openpyxl_and_its_dependency():
    wheels = _Path.home() / ".capex-tools" / "appsail-staging" / "wheels"
    if not wheels.is_dir():
        pytest.skip("provider-specific: the AppSail wheel staging directory is on the "
                    "release workstation only")
    names = [p.name for p in wheels.iterdir()]
    assert any(n.startswith("openpyxl-3.1.5") for n in names), names
    assert any(n.startswith("et_xmlfile-") for n in names), names


# =========================================================================
# Live
# =========================================================================
def _requester(pg_connection, ids) -> str:
    user = f"U-XLSX-{ids['entity']}"
    pg_connection.execute(
        "INSERT INTO app_user (user_id, email, display_name, created_by, updated_by) "
        "VALUES (%s, %s, %s, 'T', 'T')", (user, f"{user}@example.test", user))
    pg_connection.execute(
        "INSERT INTO user_access_flag (user_id, read_all, updated_by) VALUES (%s, true, 'T')",
        (user,))
    pg_connection.commit()
    return user


def _run_job(pg_database, scope, user, ids, fmt):
    with pg_database.session(scope) as session:
        job_id = export_svc.create_job(
            session, dataset="budget_ledger_cells",
            filters={"project_ids": [ids["project"]]}, scope=scope,
            requested_by=user, chunk_rows=2, output_format=fmt)["export_job_id"]
    for _ in range(30):
        if not export_svc.advance_job(pg_database, job_id, rows_budget=2).get("more"):
            break
    with pg_database.session(scope) as session:
        return job_id, export_svc.read_result(session, job_id, requester=user)


@pytest.mark.pg
@PG
def test_an_xlsx_job_runs_the_same_chunked_pipeline_and_serves_a_workbook(pg_database, pg_connection):
    ids = _seed(pg_connection, suffix=uuid.uuid4().hex[:10])
    user = _requester(pg_connection, ids)
    scope = principal_scope.scope_for_request(pg_database, {"user_id": user})

    job_id, (body, meta) = _run_job(pg_database, scope, user, ids, "xlsx")
    assert isinstance(body, (bytes, bytearray))
    assert meta["media_type"] == xlsx.MEDIA_TYPE
    assert re.fullmatch(r"budget_ledger_cells-\d{4}-\d{2}-\d{2}\.xlsx", meta["filename"])
    assert meta["csv_sha256"] and meta["xlsx_sha256"] and meta["rows"] == 5
    wb = load_workbook(io.BytesIO(bytes(body)))
    data = wb["Data"]
    header = [c.value for c in data[1]]
    assert {"wbs_id", "budget", "exposure", "internal_allocation"} <= set(header)
    assert data.max_row == 5 + 2, "five cells, a header and a totals row"
    budget_col = header.index("budget") + 1
    assert data.cell(row=2, column=budget_col).data_type == "n"
    assert data.cell(row=2, column=budget_col).number_format == xlsx.INDIAN_CURRENCY_FORMAT
    assert data.cell(row=7, column=budget_col).value.startswith("=SUM(")
    chunks = pg_connection.execute(
        "SELECT COUNT(*), SUM(row_count) FROM export_job_chunk WHERE export_job_id = %s",
        (job_id,)).fetchone()
    assert (int(chunks[0]), int(chunks[1])) == (3, 5), "chunked exactly as a csv job is"
    stored = pg_connection.execute(
        "SELECT output_format, result_media_type FROM export_job WHERE export_job_id = %s",
        (job_id,)).fetchone()
    assert stored == ("xlsx", "text/csv; charset=utf-8"), (
        "the stored chunks are CSV; the workbook is rendered from them at download")

    _csv_id, (text, csv_meta) = _run_job(pg_database, scope, user, ids, "csv")
    assert isinstance(text, str) and csv_meta["filename"].endswith(".csv")
    assert csv_meta["sha256"] == meta["csv_sha256"], "the same rows, the same digest"


@pytest.mark.pg
@PG
def test_an_unknown_format_is_refused_at_creation(pg_database, pg_connection):
    ids = _seed(pg_connection, suffix=uuid.uuid4().hex[:10])
    user = _requester(pg_connection, ids)
    scope = principal_scope.scope_for_request(pg_database, {"user_id": user})
    with pytest.raises(export_svc.ExportError) as exc:
        with pg_database.session(scope) as session:
            export_svc.create_job(session, dataset="budget_ledger_cells", filters=None,
                                  scope=scope, requested_by=user, output_format="pdf")
    assert exc.value.code == "EXPORT_FORMAT_UNKNOWN" and exc.value.status == 422
    assert pg_connection.execute("SELECT COUNT(*) FROM export_job").fetchone()[0] == 0
